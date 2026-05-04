#!/usr/bin/env python3
"""
Combine new Excel leads with existing JSON leads, enrich, and create master file
"""
import openpyxl
import json
import os
from datetime import datetime

def read_excel_leads(excel_path):
    """Read leads from Excel file"""
    print(f"Reading Excel file: {excel_path}")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    
    # Get headers from row 1
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)
    
    print(f"Headers: {headers}")
    print(f"Total rows: {ws.max_row - 1}")
    
    leads = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        lead = {}
        for i, cell_value in enumerate(row):
            if i < len(headers) and headers[i]:
                lead[headers[i]] = cell_value
        leads.append(lead)
    
    print(f"Read {len(leads)} leads from Excel")
    return leads

def convert_excel_to_json_format(excel_leads):
    """Convert Excel leads to match existing JSON format"""
    json_leads = []
    
    for excel_lead in excel_leads:
        # Map Excel columns to JSON fields
        json_lead = {
            "input_id": f"excel-{datetime.now().strftime('%Y%m%d')}-{len(json_leads)}",
            "title": excel_lead.get("Name", ""),
            "category": excel_lead.get("Category", ""),
            "address": excel_lead.get("Address", ""),
            "phone": excel_lead.get("Phone", ""),
            "website": excel_lead.get("Website", "") if excel_lead.get("Website") else None,
            "review_rating": float(excel_lead.get("Rating", 0)) if excel_lead.get("Rating") else 0.0,
            "review_count": int(excel_lead.get("Reviews", 0)) if excel_lead.get("Reviews") else 0,
            "priority_score": int(excel_lead.get("Priority Score", 0)) if excel_lead.get("Priority Score") else 0,
            "has_website": excel_lead.get("Has Website", "False") == "True",
            "website_quality": excel_lead.get("Website Quality", ""),
            "booking_method": excel_lead.get("Booking Method", ""),
            "tech_signals": excel_lead.get("Tech Signals", ""),
            "pain_points": excel_lead.get("Pain Points", ""),
            "automation_opportunities": excel_lead.get("Automation Opportunities", ""),
            "data_quality": excel_lead.get("Data Quality", ""),
            "industry": excel_lead.get("Category", ""),
            "location": excel_lead.get("Area", ""),
            "source": "excel_2026_03_30",
            "input_source": "Lead Gen - 2026-03-30.xlsx"
        }
        json_leads.append(json_lead)
    
    return json_leads

def read_existing_leads(json_path):
    """Read existing leads from JSON file"""
    if not os.path.exists(json_path):
        print(f"Existing JSON not found: {json_path}")
        return []
    
    print(f"Reading existing leads from: {json_path}")
    leads = []
    with open(json_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    lead = json.loads(line)
                    leads.append(lead)
                except json.JSONDecodeError:
                    continue
    
    print(f"Read {len(leads)} existing leads")
    return leads

def combine_leads(existing_leads, new_leads):
    """Combine existing and new leads, removing duplicates"""
    print(f"\nCombining leads...")
    print(f"Existing: {len(existing_leads)}")
    print(f"New: {len(new_leads)}")
    
    # Create a set of existing titles/phones to check duplicates
    existing_identifiers = set()
    for lead in existing_leads:
        title = (lead.get("title", "") or "").lower().strip()
        phone = (lead.get("phone", "") or "").strip()
        if title:
            existing_identifiers.add(f"title:{title}")
        if phone:
            existing_identifiers.add(f"phone:{phone}")
    
    # Add new leads that don't duplicate existing
    combined = existing_leads.copy()
    added = 0
    skipped = 0
    
    for new_lead in new_leads:
        title = (new_lead.get("title", "") or "").lower().strip()
        phone = (new_lead.get("phone", "") or "").strip()
        
        is_duplicate = False
        if title and f"title:{title}" in existing_identifiers:
            is_duplicate = True
        if phone and f"phone:{phone}" in existing_identifiers:
            is_duplicate = True
        
        if not is_duplicate:
            combined.append(new_lead)
            added += 1
            if title:
                existing_identifiers.add(f"title:{title}")
            if phone:
                existing_identifiers.add(f"phone:{phone}")
        else:
            skipped += 1
    
    print(f"Added: {added}")
    print(f"Skipped (duplicates): {skipped}")
    print(f"Total combined: {len(combined)}")
    
    return combined

def enrich_leads(leads):
    """Add enrichment data like color coding, scoring, etc."""
    print(f"\nEnriching {len(leads)} leads...")
    
    for lead in leads:
        # Add color coding based on priority score
        priority = lead.get("priority_score", 0)
        if priority >= 8:
            lead["color_code"] = "red"  # High priority
            lead["priority_label"] = "HIGH"
        elif priority >= 5:
            lead["color_code"] = "orange"  # Medium priority
            lead["priority_label"] = "MEDIUM"
        else:
            lead["color_code"] = "green"  # Low priority
            lead["priority_label"] = "LOW"
        
        # Add website status
        has_website = lead.get("has_website", False)
        website = lead.get("website")
        if not has_website or not website:
            lead["website_status"] = "NO_WEBSITE"
        elif lead.get("website_quality") == "1":
            lead["website_status"] = "POOR_WEBSITE"
        else:
            lead["website_status"] = "HAS_WEBSITE"
        
        # Add lead quality score (combine rating, reviews, priority)
        rating = lead.get("review_rating", 0) or 0
        reviews = lead.get("review_count", 0) or 0
        priority = lead.get("priority_score", 0) or 0
        
        # Simple quality score: priority * 10 + rating * 2 + log(reviews + 1)
        import math
        quality_score = (priority * 10) + (rating * 2) + math.log1p(reviews)
        lead["quality_score"] = round(quality_score, 2)
    
    print("Enrichment complete")
    return leads

def save_master_file(leads, output_path):
    """Save combined leads to master JSON file"""
    print(f"\nSaving master file to: {output_path}")
    
    with open(output_path, 'w', encoding='utf-8') as f:
        for lead in leads:
            f.write(json.dumps(lead, ensure_ascii=False) + '\n')
    
    print(f"Saved {len(leads)} leads to master file")

def create_summary(leads):
    """Create summary statistics"""
    total = len(leads)
    with_website = sum(1 for l in leads if l.get("has_website"))
    without_website = total - with_website
    high_priority = sum(1 for l in leads if l.get("priority_score", 0) >= 8)
    medium_priority = sum(1 for l in leads if 5 <= l.get("priority_score", 0) < 8)
    low_priority = sum(1 for l in leads if l.get("priority_score", 0) < 5)
    
    # By industry
    industries = {}
    for lead in leads:
        ind = lead.get("industry", "Unknown")
        industries[ind] = industries.get(ind, 0) + 1
    
    summary = {
        "total_leads": total,
        "with_website": with_website,
        "without_website": without_website,
        "high_priority": high_priority,
        "medium_priority": medium_priority,
        "low_priority": low_priority,
        "by_industry": industries,
        "generated_on": datetime.now().isoformat()
    }
    
    return summary

def main():
    # Paths
    base_dir = r"C:\Users\joshm\Desktop\App\EMBER WEB\Website search"
    excel_path = os.path.join(base_dir, "new_leads.xlsx")
    existing_json = os.path.join(base_dir, ".tmp", "all_leads.json")
    master_json = os.path.join(base_dir, ".tmp", "master_leads.json")
    summary_path = os.path.join(base_dir, ".tmp", "leads_summary.json")
    
    # Step 1: Read Excel leads
    print("="*60)
    print("STEP 1: Reading new Excel leads")
    print("="*60)
    excel_leads = read_excel_leads(excel_path)
    
    # Step 2: Convert to JSON format
    print("\n" + "="*60)
    print("STEP 2: Converting to JSON format")
    print("="*60)
    json_leads = convert_excel_to_json_format(excel_leads)
    
    # Step 3: Read existing leads
    print("\n" + "="*60)
    print("STEP 3: Reading existing leads")
    print("="*60)
    existing_leads = read_existing_leads(existing_json)
    
    # Step 4: Combine leads
    print("\n" + "="*60)
    print("STEP 4: Combining leads")
    print("="*60)
    combined_leads = combine_leads(existing_leads, json_leads)
    
    # Step 5: Enrich leads
    print("\n" + "="*60)
    print("STEP 5: Enriching leads")
    print("="*60)
    enriched_leads = enrich_leads(combined_leads)
    
    # Step 6: Save master file
    print("\n" + "="*60)
    print("STEP 6: Saving master file")
    print("="*60)
    save_master_file(enriched_leads, master_json)
    
    # Step 7: Create and save summary
    print("\n" + "="*60)
    print("STEP 7: Creating summary")
    print("="*60)
    summary = create_summary(enriched_leads)
    with open(summary_path, 'w') as f:
        json.dump(summary, f, indent=2)
    print(f"Summary saved to: {summary_path}")
    print(f"\nSummary:")
    print(json.dumps(summary, indent=2))
    
    print("\n" + "="*60)
    print("DONE! Master file created successfully")
    print("="*60)
    
    return master_json, summary

if __name__ == "__main__":
    main()
